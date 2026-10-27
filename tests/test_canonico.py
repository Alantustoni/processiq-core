"""Sugestão de fornecedor canônico: qual grafia manter, qual CNPJ, e quando não sugerir."""

import csv
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from processiq_core.canonico import (
    abreviacoes, aplicar_unificacao, de_para, nomes_compativeis, razao_completa, sugerir_canonico,
)
from processiq_core.cli import main as cli
from processiq_core.detectores import analisar, resumo
from processiq_core.exportacao import exportar_excel_bytes
from processiq_core.modelo import TipoAlerta, Titulo
from processiq_core.normalizacao import formatar_documento

HOJE = date(2026, 5, 10)
MATRIZ = "90100200000109"
FILIAL = "90100200002586"


def t(id, forn, doc=MATRIZ, linha=2, valor="100.00", venc=date(2026, 6, 1)):
    return Titulo(id=id, fornecedor=forn, documento=doc, valor=Decimal(valor), vencimento=venc,
                  linha=linha, documento_bruto=doc)


# ---------- critérios isolados ----------

def test_abreviacoes_e_razao_completa():
    assert abreviacoes("Distrib. Alpha Ltda") == 1
    assert abreviacoes("Com. e Ind. Beta S/A") == 2
    assert abreviacoes("Distribuidora Alpha Ltda.") == 0          # LTDA. não é abreviação de nome
    assert razao_completa("Transportes Fictícios S/A") and razao_completa("Peças Exemplo Eireli")
    assert not razao_completa("Mercado Central")


def test_nomes_compativeis_por_prefixo_e_similaridade():
    assert nomes_compativeis("Merc. Central", "Mercado Central Alimentos")      # prefixos, em ordem
    assert nomes_compativeis("Distrib. Alpha", "Distribuidora Alpha Ltda")
    assert nomes_compativeis("Mercado Central Ltda", "MERCADO CENTRAL")         # iguais após normalizar
    assert not nomes_compativeis("Transportes Fictícios", "Serviços Demonstração")
    assert not nomes_compativeis("Central", "Central de Peças Norte")           # uma palavra só não basta


# ---------- escolha da grafia ----------

def test_extenso_vence_abreviado_mesmo_sendo_menos_usado():
    tit = [t("1", "Distrib. Alpha Ltda", linha=2), t("2", "Distrib. Alpha Ltda", linha=3),
           t("3", "Distrib. Alpha Ltda", linha=4), t("4", "Distribuidora Alpha Ltda", linha=5)]
    s = sugerir_canonico(tit)
    assert s.nome == "Distribuidora Alpha Ltda"
    assert "por extenso" in s.motivo
    assert {v.nome for v in s.variantes} == {"Distrib. Alpha Ltda", "Distribuidora Alpha Ltda"}


def test_grafia_mais_usada_vence():
    tit = [t("1", "Mercado Central Ltda", linha=2), t("2", "Mercado Central Ltda", linha=3),
           t("3", "Mercado Central Ltda", linha=4), t("4", "Mercado Central Alimentos Ltda", linha=5)]
    s = sugerir_canonico(tit)
    assert s.nome == "Mercado Central Ltda"
    assert "3 de 4" in s.motivo


def test_razao_social_completa_desempata():
    tit = [t("1", "Beta Tintas Industrial", linha=2), t("2", "Beta Tintas Ltda", linha=3)]
    s = sugerir_canonico(tit)
    assert s.nome == "Beta Tintas Ltda"
    assert "razão social completa" in s.motivo


def test_mesmo_nome_com_grafias_diferentes_prefere_acento_e_tipo_societario():
    tit = [t("1", "ALPHA PECAS", linha=2), t("2", "Alpha Peças Ltda", linha=3)]
    s = sugerir_canonico(tit)
    assert s.nome == "Alpha Peças Ltda"


def test_cnpj_sugerido_e_o_da_matriz():
    tit = [t("1", "Mercado Central", doc=FILIAL, linha=2), t("2", "Mercado Central", doc=FILIAL, linha=3),
           t("3", "Merc. Central", doc=MATRIZ, linha=4)]
    assert sugerir_canonico(tit).documento == MATRIZ
    assert formatar_documento(MATRIZ) == "90.100.200/0001-09"


def test_nomes_sem_relacao_nao_viram_sugestao():
    tit = [t("1", "Transportes Fictícios S/A", linha=2), t("2", "Serviços Demonstração Ltda", linha=3)]
    s = sugerir_canonico(tit)
    assert s.nome == "" and s.variantes == ()
    assert {v.nome for v in s.conferir} == {"Transportes Fictícios S/A", "Serviços Demonstração Ltda"}
    assert "CNPJ digitado errado" in s.motivo


def test_grafia_sem_relacao_sai_para_conferir_e_o_resto_e_unificado():
    tit = [t("1", "Mercado Central Ltda", linha=2), t("2", "MERC CENTRAL", linha=3),
           t("3", "Oficina Zeta", linha=4)]                     # mesmo CNPJ, nada a ver com o resto
    s = sugerir_canonico(tit)
    assert s.nome == "Mercado Central Ltda"
    assert [v.nome for v in s.conferir] == ["Oficina Zeta"]


# ---------- integração com o detector ----------

def test_alerta_traz_sugestao_na_mensagem_e_nos_detalhes():
    tit = [t("1", "Mercado Central Ltda", linha=2), t("2", "Merc. Central Alimentos", doc=FILIAL, linha=3)]
    al = [a for a in analisar(tit, hoje=HOJE) if a.tipo == TipoAlerta.FORNECEDOR_DUPLICADO]
    assert len(al) == 1
    assert "manter 'Mercado Central Ltda'" in al[0].mensagem
    c = al[0].detalhes["canonico"]
    assert c["nome"] == "Mercado Central Ltda" and c["documento"] == MATRIZ


def test_de_para_separa_unificar_padronizar_e_conferir():
    tit = [t("1", "Mercado Central Ltda", linha=2), t("2", "Mercado Central Ltda", linha=3),
           t("3", "MERCADO CENTRAL", doc=FILIAL, linha=4), t("4", "Merc. Central", linha=5),
           t("5", "Oficina Zeta", linha=6)]
    tabela = de_para(analisar(tit, hoje=HOJE))
    por_de = {r["de"]: r for r in tabela}
    assert por_de["MERCADO CENTRAL"]["acao"] == "padronizar grafia"
    assert por_de["Merc. Central"]["acao"] == "unificar"
    assert por_de["Merc. Central"]["para"] == "Mercado Central Ltda" and por_de["Merc. Central"]["linhas"] == [5]
    assert por_de["Oficina Zeta"]["acao"] == "conferir" and por_de["Oficina Zeta"]["para"] == ""
    assert "Mercado Central Ltda" not in por_de                  # a grafia mantida não aparece como troca


def test_aplicar_unificacao_revela_nf_duplicada_escondida():
    # a mesma cobrança lançada com duas grafias: antes de unificar, o detector de NF duplicada
    # não pega (fornecedores diferentes); depois, pega.
    tit = [t("1", "Mercado Central Ltda", linha=2, valor="4820.00"),
           t("2", "Merc. Central Alimentos", doc=FILIAL, linha=3, valor="4820.00")]
    antes = analisar(tit, hoje=HOJE)
    assert not any(a.tipo == TipoAlerta.TITULO_DUPLICADO for a in antes)

    limpos = aplicar_unificacao(tit, de_para(antes))
    assert [x.fornecedor for x in limpos] == ["Mercado Central Ltda", "Mercado Central Ltda"]
    depois = analisar(limpos, hoje=HOJE)
    assert not any(a.tipo == TipoAlerta.FORNECEDOR_DUPLICADO for a in depois)
    assert any(a.tipo == TipoAlerta.TITULO_DUPLICADO for a in depois)


def test_aplicar_unificacao_nao_mexe_em_casos_para_conferir():
    tit = [t("1", "Transportes Fictícios S/A", linha=2), t("2", "Serviços Demonstração Ltda", linha=3)]
    assert aplicar_unificacao(tit, de_para(analisar(tit, hoje=HOJE))) == tit


# ---------- saídas: Excel, CSV ----------

def test_aba_unificacao_no_excel():
    tit = [t("1", "Mercado Central Ltda", linha=2), t("2", "Merc. Central", doc=FILIAL, linha=3)]
    al = analisar(tit, hoje=HOJE)
    wb = load_workbook(BytesIO(exportar_excel_bytes(tit, al, resumo(tit, al, hoje=HOJE), HOJE)))
    ws = wb["Unificação"]
    assert [c.value for c in ws[1]][:4] == ["Ação", "De (grafia atual)", "Para (manter)", "CNPJ sugerido"]
    assert [c.value for c in ws[2]][:4] == ["unificar", "Merc. Central", "Mercado Central Ltda", "90.100.200/0001-09"]


def test_cli_grava_de_para_em_csv(tmp_path, capsys):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["NF", "Fornecedor", "CNPJ", "Valor", "Vencimento"])
    ws.append([1, "Distribuidora Alpha Ltda", "90.100.200/0001-09", 100, date(2026, 6, 1)])
    ws.append([2, "Distrib. Alpha", "90.100.200/0025-86", 200, date(2026, 6, 2)])
    origem, destino = tmp_path / "cap.xlsx", tmp_path / "unificacao.csv"
    wb.save(origem)
    assert cli([str(origem), "--de-para", str(destino), "--hoje", "2026-05-10"]) == 0
    with open(destino, encoding="utf-8-sig", newline="") as f:
        linhas = list(csv.reader(f, delimiter=";"))
    assert linhas[0][:4] == ["acao", "de", "para", "cnpj_sugerido"]
    assert linhas[1][:4] == ["unificar", "Distrib. Alpha", "Distribuidora Alpha Ltda", "90.100.200/0001-09"]
    assert "manter 'Distribuidora Alpha Ltda'" in capsys.readouterr().out
