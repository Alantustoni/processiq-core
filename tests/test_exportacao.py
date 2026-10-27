"""Gera o relatório Excel e lê de volta com openpyxl para conferir conteúdo e marcação."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from processiq_core.detectores import analisar, resumo
from processiq_core.exportacao import CORES, exportar_excel, exportar_excel_bytes
from processiq_core.modelo import Severidade, Titulo

HOJE = date(2026, 5, 10)
CNPJ = "90100200000109"


def _t(id, forn, doc=CNPJ, valor="100.00", venc=date(2026, 5, 20), linha=2, status=""):
    return Titulo(id=id, fornecedor=forn, documento=doc, valor=Decimal(valor), vencimento=venc,
                  linha=linha, status=status, documento_bruto=doc)


def _cenario():
    titulos = [
        _t("1001", "Mercado Central", valor="4820.00", venc=date(2026, 5, 12), linha=2),
        _t("1001", "MERCADO CENTRAL LTDA", valor="4820.00", venc=date(2026, 5, 12), linha=3),   # duplicada → crítica
        _t("1002", "Distribuidora Alpha", doc="90100300001700", venc=date(2026, 5, 1), linha=4),  # vencida → alta
        _t("1003", "Clinica Sao Lucas", doc="90100400000838", venc=date(2026, 7, 1), linha=5),    # limpa
    ]
    alertas = analisar(titulos, hoje=HOJE)
    return titulos, alertas, resumo(titulos, alertas, hoje=HOJE)


def test_gera_tres_abas_e_marca_linhas(tmp_path: Path):
    titulos, alertas, res = _cenario()
    destino = exportar_excel(titulos, alertas, res, HOJE, tmp_path / "rel.xlsx", origem="cap.xlsx")
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Títulos", "Alertas", "Resumo"]

    ws = wb["Títulos"]
    cab = [c.value for c in ws[1]]
    assert cab[:4] == ["Linha", "ID", "Fornecedor", "CNPJ/CPF"] and cab[-2:] == ["Alertas", "Severidade"]
    linhas = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    assert set(linhas) == {2, 3, 4, 5}

    # linha 2 tem NF duplicada (crítica) e vence em 2 dias (baixa): fica com a pior cor
    r = linhas[2]
    assert ws.cell(r, 10).value == "critica"
    assert ws.cell(r, 1).fill.fgColor.rgb.endswith(CORES[Severidade.CRITICA])
    assert "lançamentos iguais" in ws.cell(r, 9).value

    # linha 4 vencida → alta
    r = linhas[4]
    assert ws.cell(r, 10).value == "alta"
    assert ws.cell(r, 1).fill.fgColor.rgb.endswith(CORES[Severidade.ALTA])

    # linha 5 sem alerta: sem severidade e sem preenchimento sólido
    r = linhas[5]
    assert ws.cell(r, 10).value in ("", None)
    assert ws.cell(r, 1).fill.fill_type in (None, "none")


def test_aba_alertas_e_resumo(tmp_path: Path):
    titulos, alertas, res = _cenario()
    wb = load_workbook(exportar_excel(titulos, alertas, res, HOJE, tmp_path / "rel.xlsx"))
    wa = wb["Alertas"]
    assert wa.max_row - 1 == len(alertas)
    assert [c.value for c in wa[1]] == ["Severidade", "Tipo", "Mensagem", "Linhas", "Títulos"]
    assert wa.cell(2, 1).value == "critica"                      # ordenado por severidade

    wr = wb["Resumo"]
    valores = {wr.cell(r, 1).value: wr.cell(r, 2).value for r in range(1, wr.max_row + 1)}
    assert valores["Títulos"] == 4
    assert valores["Valor total"] == float(Decimal(res["valor_total"]))
    assert valores["Data de referência"] == "10/05/2026"
    assert valores["critica"] == res["por_severidade"]["critica"]


def test_exportar_bytes_abre_como_xlsx():
    titulos, alertas, res = _cenario()
    dados = exportar_excel_bytes(titulos, alertas, res, HOJE)
    assert dados[:2] == b"PK"                                    # zip → xlsx
    wb = load_workbook(BytesIO(dados))
    assert wb["Títulos"].max_row == 5


def test_sem_alertas_gera_relatorio_limpo(tmp_path: Path):
    titulos = [_t("1", "Fornecedor Unico", venc=date(2026, 7, 1))]
    alertas = analisar(titulos, hoje=HOJE)
    assert alertas == []
    wb = load_workbook(exportar_excel(titulos, alertas, resumo(titulos, alertas, hoje=HOJE), HOJE, tmp_path / "r.xlsx"))
    assert wb["Alertas"].max_row == 1                             # só cabeçalho
    assert wb["Títulos"].cell(2, 10).value in ("", None)
