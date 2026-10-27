"""Testes de integração da API: sobem a app em memória (TestClient) e enviam planilhas reais."""

from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from processiq_core.api import app

client = TestClient(app)
CNPJ = "90.100.200/0001-09"


def _xlsx(linhas: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for l in linhas:
        ws.append(l)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _planilha_padrao() -> bytes:
    return _xlsx([
        ["NF", "Fornecedor", "CNPJ", "Valor", "Vencimento", "Status"],
        [1001, "Mercado Central Ltda", CNPJ, 4820.00, date(2026, 5, 12), "Pendente"],
        [1001, "MERCADO CENTRAL LTDA", CNPJ, 4820.00, date(2026, 5, 12), "Pendente"],   # NF duplicada
        [1002, "Distribuidora Alpha", "", 980.00, date(2026, 5, 1), "Pendente"],          # vencida
        [1003, "Clínica São Lucas", "12.345.678/0001-00", 300.00, date(2026, 6, 1), ""],  # CNPJ inválido
    ])


def test_health():
    r = client.get("/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_analisar_xlsx_devolve_colunas_resumo_e_alertas():
    r = client.post("/analisar", params={"hoje": "2026-05-10"},
                    files={"arquivo": ("cap.xlsx", _planilha_padrao(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["arquivo"] == "cap.xlsx"
    assert body["hoje"] == "2026-05-10"
    assert body["colunas"]["fornecedor"] == 1 and body["colunas"]["vencimento"] == 4
    assert body["resumo"]["titulos"] == 4
    tipos = {a["tipo"] for a in body["alertas"]}
    assert {"titulo_duplicado", "vencido", "documento_invalido"} <= tipos
    assert body["titulos"] is None                       # só com incluir_titulos=true


def test_alertas_vem_ordenados_por_severidade():
    r = client.post("/analisar", params={"hoje": "2026-05-10"},
                    files={"arquivo": ("cap.xlsx", _planilha_padrao(), "application/octet-stream")})
    ordem = {"critica": 0, "alta": 1, "media": 2, "baixa": 3}
    sev = [ordem[a["severidade"]] for a in r.json()["alertas"]]
    assert sev == sorted(sev)


def test_incluir_titulos():
    r = client.post("/analisar", params={"hoje": "2026-05-10", "incluir_titulos": "true"},
                    files={"arquivo": ("cap.xlsx", _planilha_padrao(), "application/octet-stream")})
    t = r.json()["titulos"]
    assert len(t) == 4
    assert t[0]["documento"] == "90100200000109" and t[0]["valor"] == "4820.00" and t[0]["linha"] == 2


def test_csv_latin1_ponto_e_virgula():
    csv = "Fornecedor;CNPJ;Valor;Vencimento\nDistribuidora Alpha;90.100.300/0017-00;1.234,56;03/05/2026\n".encode("latin-1")
    r = client.post("/analisar", params={"hoje": "2026-05-10"},
                    files={"arquivo": ("cap.csv", csv, "text/csv")})
    assert r.status_code == 200, r.text
    assert r.json()["resumo"]["valor_total"] == "1234.56"


def test_formato_nao_suportado_415():
    r = client.post("/analisar", files={"arquivo": ("x.pdf", b"%PDF", "application/pdf")})
    assert r.status_code == 415
    assert "formato" in r.json()["detail"]


def test_arquivo_vazio_422():
    r = client.post("/analisar", files={"arquivo": ("x.csv", b"", "text/csv")})
    assert r.status_code == 422


def test_cabecalho_nao_reconhecido_422_com_mensagem_do_motor():
    r = client.post("/analisar", files={"arquivo": ("x.xlsx", _xlsx([["a", "b"], [1, 2]]), "application/octet-stream")})
    assert r.status_code == 422
    assert "cabeçalho não reconhecido" in r.json()["erro"]


def test_arquivo_grande_413(monkeypatch):
    import sys
    modulo = sys.modules["processiq_core.api.app"]      # o atributo `app` do pacote esconde o submódulo
    monkeypatch.setattr(modulo, "TAMANHO_MAXIMO", 100)
    r = client.post("/analisar", files={"arquivo": ("cap.xlsx", _planilha_padrao(), "application/octet-stream")})
    assert r.status_code == 413


def test_openapi_documenta_o_endpoint():
    spec = client.get("/openapi.json").json()
    assert "/analisar" in spec["paths"]
    assert "AnaliseOut" in spec["components"]["schemas"]


def test_analisar_excel_devolve_xlsx_com_linhas_marcadas():
    from openpyxl import load_workbook
    r = client.post("/analisar/excel", params={"hoje": "2026-05-10"},
                    files={"arquivo": ("cap.xlsx", _planilha_padrao(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="cap_processiq.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Títulos", "Alertas", "Unificação", "Resumo"]
    ws = wb["Títulos"]
    sev = {ws.cell(i, 1).value: ws.cell(i, 10).value for i in range(2, ws.max_row + 1)}
    assert sev[2] == "critica" and sev[3] == "critica"           # NF duplicada nas linhas 2 e 3


def test_analisar_excel_formato_invalido_415():
    r = client.post("/analisar/excel", files={"arquivo": ("x.pdf", b"%PDF", "application/pdf")})
    assert r.status_code == 415


def test_analisar_devolve_tabela_de_unificacao():
    planilha = _xlsx([
        ["NF", "Fornecedor", "CNPJ", "Valor", "Vencimento"],
        [1, "Distribuidora Alpha Ltda", CNPJ, 100.0, date(2026, 6, 1)],
        [2, "Distrib. Alpha", "90.100.200/0025-86", 200.0, date(2026, 6, 2)],
    ])
    r = client.post("/analisar", params={"hoje": "2026-05-10"},
                    files={"arquivo": ("cap.xlsx", planilha, "application/octet-stream")})
    assert r.status_code == 200, r.text
    u = r.json()["unificacao"]
    assert u == [{"acao": "unificar", "de": "Distrib. Alpha", "para": "Distribuidora Alpha Ltda",
                  "documento": "90100200000109", "ocorrencias": 1, "linhas": [3], "criterio": "cnpj",
                  "motivo": u[0]["motivo"]}]
    assert "por extenso" in u[0]["motivo"]
