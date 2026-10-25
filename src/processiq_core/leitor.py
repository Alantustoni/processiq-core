"""Leitura de planilhas de contas a pagar (.xlsx/.csv) com detecção de colunas pelo cabeçalho.

Pequenas empresas não têm layout padrão: a coluna do CNPJ pode se chamar
'CNPJ', 'Documento', 'CPF/CNPJ'; o vencimento pode ser 'Vencimento', 'Data Venc.',
'Dt. Vencto'. Em vez de exigir um modelo, procuramos cada campo por sinônimos
e devolvemos quais colunas foram reconhecidas, para o usuário conferir.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Iterable

from .modelo import Titulo
from .normalizacao import normalizar_documento, normalizar_nome, parse_data, parse_valor, somente_digitos

# campo lógico → sinônimos de cabeçalho (já normalizados por normalizar_nome)
SINONIMOS: dict[str, tuple[str, ...]] = {
    "id":          ("ID", "NF", "NO NF", "N NF", "NUM NF", "NUMERO NF", "NOTA", "NOTA FISCAL", "NUMERO", "N", "NUM", "DOCUMENTO FISCAL", "TITULO", "CODIGO"),
    "fornecedor":  ("FORNECEDOR", "RAZAO SOCIAL", "NOME", "FAVORECIDO", "BENEFICIARIO", "CREDOR", "EMPRESA"),
    "documento":   ("CNPJ", "CPF", "CPF CNPJ", "CNPJ CPF", "DOCUMENTO", "DOC", "CNPJ FORNECEDOR"),
    "valor":       ("VALOR", "VALOR TOTAL", "TOTAL", "VLR", "MONTANTE", "VALOR NF", "VALOR TITULO"),
    "vencimento":  ("VENCIMENTO", "DATA VENCIMENTO", "DT VENCIMENTO", "VENC", "VENCTO", "DT VENC", "DATA VENC"),
    "categoria":   ("CATEGORIA", "TIPO", "CENTRO CUSTO", "NATUREZA", "CLASSIFICACAO"),
    "status":      ("STATUS", "SITUACAO", "PAGO", "ESTADO"),
}

CAMPOS_OBRIGATORIOS = ("fornecedor", "valor", "vencimento")


class ErroLeitura(ValueError):
    pass


def mapear_colunas(cabecalho: Iterable) -> dict[str, int]:
    """Devolve {campo_logico: indice_da_coluna} para os campos reconhecidos.

    Preferência: correspondência exata com um sinônimo; depois, cabeçalho que
    contenha o sinônimo como palavra. 'documento' só é aceito se não bater
    antes com 'id' (evita 'Documento Fiscal' virar CNPJ).
    """
    nomes = [normalizar_nome(str(c) if c is not None else "") for c in cabecalho]
    mapa: dict[str, int] = {}
    usados: set[int] = set()

    for campo, sins in SINONIMOS.items():                      # 1ª passada: exato
        for i, nome in enumerate(nomes):
            if i in usados or not nome:
                continue
            if nome in sins:
                mapa[campo] = i
                usados.add(i)
                break

    for campo, sins in SINONIMOS.items():                      # 2ª passada: contém
        if campo in mapa:
            continue
        for i, nome in enumerate(nomes):
            if i in usados or not nome:
                continue
            palavras = set(nome.split())
            if any(set(s.split()) <= palavras for s in sins if len(s) > 2):
                mapa[campo] = i
                usados.add(i)
                break
    return mapa


def _linhas_xlsx(caminho: Path) -> list[list]:
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _linhas_csv(caminho: Path) -> list[list]:
    raw = caminho.read_bytes()
    for enc in ("utf-8-sig", "latin-1"):
        try:
            texto = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErroLeitura("não foi possível decodificar o CSV (utf-8 ou latin-1)")
    amostra = texto[:4096]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
    except csv.Error:
        dialeto = csv.excel
        dialeto.delimiter = ";" if amostra.count(";") >= amostra.count(",") else ","
    return [row for row in csv.reader(io.StringIO(texto), dialeto)]


def _achar_cabecalho(linhas: list[list], max_busca: int = 15) -> int:
    """Planilhas de pequena empresa têm título/logo nas primeiras linhas; procuramos
    a primeira linha em que pelo menos os campos obrigatórios são reconhecidos."""
    for i, linha in enumerate(linhas[:max_busca]):
        mapa = mapear_colunas(linha)
        if all(c in mapa for c in CAMPOS_OBRIGATORIOS):
            return i
    raise ErroLeitura(
        "cabeçalho não reconhecido: preciso de colunas de fornecedor, valor e vencimento "
        f"(procurei nas {max_busca} primeiras linhas)"
    )


def _celula(linha: list, idx: int | None):
    if idx is None or idx >= len(linha):
        return None
    return linha[idx]


def ler_planilha(caminho: str | Path) -> tuple[list[Titulo], dict[str, int]]:
    """Lê .xlsx ou .csv e devolve (títulos, mapa_de_colunas)."""
    caminho = Path(caminho)
    if not caminho.exists():
        raise ErroLeitura(f"arquivo não encontrado: {caminho}")
    suf = caminho.suffix.lower()
    if suf in (".xlsx", ".xlsm"):
        linhas = _linhas_xlsx(caminho)
    elif suf in (".csv", ".txt"):
        linhas = _linhas_csv(caminho)
    else:
        raise ErroLeitura(f"formato não suportado: {suf} (use .xlsx ou .csv)")
    return titulos_de_linhas(linhas)


def titulos_de_linhas(linhas: list[list]) -> tuple[list[Titulo], dict[str, int]]:
    """Mesma lógica de ler_planilha, mas a partir de linhas já em memória (útil em testes e APIs)."""
    if not linhas:
        raise ErroLeitura("planilha vazia")
    h = _achar_cabecalho(linhas)
    mapa = mapear_colunas(linhas[h])
    titulos: list[Titulo] = []
    for n, linha in enumerate(linhas[h + 1:], start=h + 2):   # +2: 1-based e pula cabeçalho
        if not linha or all(c in (None, "") for c in linha):
            continue
        fornecedor = _celula(linha, mapa.get("fornecedor"))
        valor = _celula(linha, mapa.get("valor"))
        venc = _celula(linha, mapa.get("vencimento"))
        if fornecedor in (None, "") and parse_valor(valor) is None and parse_data(venc) is None:
            continue                                            # linha de total/rodapé (ex.: '=SUM(...)')
        if fornecedor in (None, "") and parse_data(venc) is None:
            continue                                            # total numérico sem fornecedor nem data
        doc_bruto = _celula(linha, mapa.get("documento"))
        ident = _celula(linha, mapa.get("id"))
        ident = str(ident).strip() if ident not in (None, "") else f"L{n}"
        if ident.endswith(".0"):
            ident = ident[:-2]
        titulos.append(Titulo(
            id=ident,
            fornecedor=str(fornecedor).strip() if fornecedor is not None else "",
            documento=normalizar_documento(doc_bruto),
            documento_bruto=somente_digitos(doc_bruto),
            valor=parse_valor(valor),
            vencimento=parse_data(venc),
            linha=n,
            categoria=str(_celula(linha, mapa.get("categoria")) or "").strip(),
            status=str(_celula(linha, mapa.get("status")) or "").strip(),
        ))
    return titulos, mapa
