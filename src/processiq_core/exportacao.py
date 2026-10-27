"""Exporta o resultado da análise para um .xlsx que o financeiro abre e já sabe onde olhar.

Três abas:
  • Títulos  — um título por linha, com as colunas originais normalizadas mais
               `Alertas` e `Severidade`; a linha inteira recebe a cor da pior severidade.
  • Alertas  — um alerta por linha (tipo, severidade, mensagem, linhas da planilha de origem).
  • Resumo   — indicadores do `resumo()` e legenda de cores.

Usa só openpyxl (já é dependência do leitor). Nenhuma regra de negócio aqui:
este módulo recebe títulos e alertas prontos e só os desenha.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .modelo import Alerta, Severidade, Titulo

# cor de preenchimento por severidade (tons claros, legíveis em impressão)
CORES = {
    Severidade.CRITICA: "F8CBAD",   # laranja-avermelhado
    Severidade.ALTA:    "FFE699",   # amarelo
    Severidade.MEDIA:   "DDEBF7",   # azul-claro
    Severidade.BAIXA:   "E2EFDA",   # verde-claro
}
_ORDEM = {Severidade.CRITICA: 0, Severidade.ALTA: 1, Severidade.MEDIA: 2, Severidade.BAIXA: 3}
_CABECALHO = PatternFill("solid", fgColor="1F4E78")
_FONTE_CAB = Font(bold=True, color="FFFFFF")


def _pior(sevs: list[Severidade]) -> Severidade | None:
    return min(sevs, key=_ORDEM.get) if sevs else None


def _alertas_por_linha(alertas: list[Alerta]) -> dict[int, list[Alerta]]:
    por_linha: dict[int, list[Alerta]] = defaultdict(list)
    for a in alertas:
        for ln in a.linhas:
            por_linha[ln].append(a)
    return por_linha


def _cabecalho(ws: Worksheet, titulos: list[str]) -> None:
    ws.append(titulos)
    for c in ws[1]:
        c.fill = _CABECALHO
        c.font = _FONTE_CAB
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _ajustar_larguras(ws: Worksheet, maximo: int = 60) -> None:
    for col in ws.columns:
        letra = get_column_letter(col[0].column)
        maior = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letra].width = min(max(10, maior + 2), maximo)


def _aba_titulos(wb: Workbook, titulos: list[Titulo], alertas: list[Alerta]) -> None:
    ws = wb.active
    ws.title = "Títulos"
    _cabecalho(ws, ["Linha", "ID", "Fornecedor", "CNPJ/CPF", "Valor", "Vencimento",
                    "Categoria", "Status", "Alertas", "Severidade"])
    por_linha = _alertas_por_linha(alertas)
    for t in titulos:
        meus = por_linha.get(t.linha, [])
        pior = _pior([a.severidade for a in meus])
        ws.append([
            t.linha, t.id, t.fornecedor, t.documento or t.documento_bruto,
            float(t.valor) if t.valor is not None else None,
            t.vencimento, t.categoria, t.status,
            "; ".join(a.mensagem for a in meus) or "",
            pior.value if pior else "",
        ])
        r = ws.max_row
        ws.cell(r, 5).number_format = '#,##0.00'
        ws.cell(r, 6).number_format = 'DD/MM/YYYY'
        if pior:
            fill = PatternFill("solid", fgColor=CORES[pior])
            for c in ws[r]:
                c.fill = fill
    _ajustar_larguras(ws)
    ws.column_dimensions["I"].width = 80


def _aba_alertas(wb: Workbook, alertas: list[Alerta]) -> None:
    ws = wb.create_sheet("Alertas")
    _cabecalho(ws, ["Severidade", "Tipo", "Mensagem", "Linhas", "Títulos"])
    for a in alertas:
        ws.append([a.severidade.value, a.tipo.value, a.mensagem,
                   ", ".join(str(l) for l in a.linhas), ", ".join(a.titulos)])
        fill = PatternFill("solid", fgColor=CORES[a.severidade])
        ws.cell(ws.max_row, 1).fill = fill
    _ajustar_larguras(ws)
    ws.column_dimensions["C"].width = 90


def _aba_resumo(wb: Workbook, resumo: dict, hoje: date, origem: str) -> None:
    ws = wb.create_sheet("Resumo")
    ws["A1"] = "ProcessIQ — análise de contas a pagar"
    ws["A1"].font = Font(bold=True, size=14)
    linhas = [
        ("Arquivo analisado", origem),
        ("Data de referência", hoje.strftime("%d/%m/%Y")),
        ("Títulos", resumo["titulos"]),
        ("Valor total", float(Decimal(resumo["valor_total"]))),
        ("Valor vencido", float(Decimal(resumo["valor_vencido"]))),
        ("Vence nos próximos 7 dias", float(Decimal(resumo["valor_proximos_7_dias"]))),
        ("Alertas", resumo["alertas"]),
        ("Saúde operacional (0–100)", resumo["saude_operacional"]),
    ]
    for i, (k, v) in enumerate(linhas, start=3):
        ws.cell(i, 1, k).font = Font(bold=True)
        c = ws.cell(i, 2, v)
        if isinstance(v, float):
            c.number_format = '#,##0.00'

    r = len(linhas) + 5
    ws.cell(r, 1, "Alertas por severidade").font = Font(bold=True)
    for sev in (Severidade.CRITICA, Severidade.ALTA, Severidade.MEDIA, Severidade.BAIXA):
        r += 1
        ws.cell(r, 1, sev.value).fill = PatternFill("solid", fgColor=CORES[sev])
        ws.cell(r, 2, resumo["por_severidade"].get(sev.value, 0))

    r += 2
    ws.cell(r, 1, "Alertas por tipo").font = Font(bold=True)
    for tipo, n in sorted(resumo["por_tipo"].items(), key=lambda kv: -kv[1]):
        r += 1
        ws.cell(r, 1, tipo)
        ws.cell(r, 2, n)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 40


def gerar_workbook(titulos: list[Titulo], alertas: list[Alerta], resumo: dict,
                   hoje: date, origem: str = "") -> Workbook:
    """Monta o workbook em memória. `resumo` é o dict de `detectores.resumo()`."""
    wb = Workbook()
    _aba_titulos(wb, titulos, alertas)
    _aba_alertas(wb, alertas)
    _aba_resumo(wb, resumo, hoje, origem)
    return wb


def exportar_excel(titulos: list[Titulo], alertas: list[Alerta], resumo: dict,
                   hoje: date, destino: str | Path, origem: str = "") -> Path:
    """Grava o relatório em `destino` e devolve o caminho."""
    destino = Path(destino)
    gerar_workbook(titulos, alertas, resumo, hoje, origem).save(destino)
    return destino


def exportar_excel_bytes(titulos: list[Titulo], alertas: list[Alerta], resumo: dict,
                         hoje: date, origem: str = "") -> bytes:
    """Mesmo relatório, como bytes — para a API devolver sem tocar em disco."""
    buf = BytesIO()
    gerar_workbook(titulos, alertas, resumo, hoje, origem).save(buf)
    return buf.getvalue()
